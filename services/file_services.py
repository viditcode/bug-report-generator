import openpyxl
import os

header = ["Bug_Id","Bug_Summary","Bug_Description","Severity","Priority", "Environment","Bug_Label", "Step_to_reproduce","Date and Time"]
class FileServices:
    def __init__(self):
        # for excel
        if os.path.exists("Bug_report.xlsx"):
            self.wb = openpyxl.load_workbook("Bug_report.xlsx")
            self.ws = self.wb.active
            current_header = [cell.value for cell in self.ws[1]]
            if current_header != header:
                for col, header_value in enumerate(header, 1):
                    self.ws.cell(row=1, column=col, value= header_value)
        else:
            self.wb = openpyxl.Workbook()
            self.ws = self.wb.active
            self.ws.append(header)

        # for text file
        if os.path.exists("Bug_report.txt"):
            self.file = open("Bug_report.txt","a") 
        else:
            self.file = open("Bug_report.txt", "a")
            self.file.write("="*40+"\n")
            self.file.write("        Bug Report"+"\n")
            self.file.write("="*40+"\n")
    
    def get_bug_number(self):
        if self.ws.max_row <= 1:
            return 1
        return self.ws.max_row
        
        
    def save_to_txt(self,bug):
        self.file.write("Bug ID: - "+bug.bug_id+"\n")
        self.file.write("Bug summary: - "+bug.summary+"\n")
        self.file.write("Bug Description: - "+bug.description+"\n")
        self.file.write("Severity: - "+bug.severity+"\n")
        self.file.write("Priority: - "+bug.priority+"\n")
        self.file.write("Environment: - "+bug.environment+"\n")
        self.file.write("Bug Type: - "+ bug.label+"\n")
        self.file.write("Steps to reproduce: - \n")
        for step_num in bug.steps:
            self.file.write(step_num+"\n")
        self.file.write("Date and Time: - "+bug.timestamp+"\n")
        self.file.write("="*40+"\n")
    
    def save_to_excel(self,bug):
        self.ws.append([
            bug.bug_id,bug.summary,bug.description,bug.severity,bug.priority,bug.environment,bug.label,",".join(bug.steps), bug.timestamp
            ])
        
    def close (self):
        self.wb.save("Bug_report.xlsx")
        self.file.close()